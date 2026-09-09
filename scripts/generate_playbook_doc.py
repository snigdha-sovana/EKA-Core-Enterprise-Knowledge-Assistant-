"""Generate Word (.docx) and styled HTML versions of INTERVIEW_DEMO_PLAYBOOK.md."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Inches, Pt, RGBColor


def create_docx(md_path: Path, output_path: Path):
    doc = Document()

    # Document margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Styles
    style_normal = doc.styles["Normal"]
    font = style_normal.font
    font.name = "Segoe UI"
    font.size = Pt(10.5)
    font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)  # Slate 800

    text = md_path.read_text(encoding="utf-8")
    lines = text.splitlines()

    i = 0
    in_code_block = False
    code_lines = []

    while i < len(lines):
        line = lines[i]

        # Handle fenced code blocks
        if line.startswith("```"):
            if in_code_block:
                # Flush code block
                code_text = "\n".join(code_lines)
                table = doc.add_table(rows=1, cols=1)
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                cell = table.cell(0, 0)
                shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F1F5F9"/>')
                cell._tc.get_or_add_tcPr().append(shading_elm)
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(4)
                run = p.add_run(code_text)
                run.font.name = "Consolas"
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
                in_code_block = False
                code_lines = []
            else:
                in_code_block = True
                code_lines = []
            i += 1
            continue

        if in_code_block:
            code_lines.append(line)
            i += 1
            continue

        stripped = line.strip()

        # Headings
        if stripped.startswith("# "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(16)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(stripped[2:])
            run.font.name = "Segoe UI"
            run.font.size = Pt(22)
            run.bold = True
            run.font.color.rgb = RGBColor(0x1E, 0x40, 0xAF)  # Royal Blue
        elif stripped.startswith("## "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(stripped[3:])
            run.font.name = "Segoe UI"
            run.font.size = Pt(15)
            run.bold = True
            run.font.color.rgb = RGBColor(0x25, 0x63, 0xEB)
        elif stripped.startswith("### "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(stripped[4:])
            run.font.name = "Segoe UI"
            run.font.size = Pt(12)
            run.bold = True
            run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
        elif stripped.startswith("#### "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(stripped[5:])
            run.font.name = "Segoe UI"
            run.font.size = Pt(11)
            run.bold = True
            run.font.color.rgb = RGBColor(0x47, 0x55, 0x69)
        elif stripped.startswith("> "):
            # Callout / Quote block
            table = doc.add_table(rows=1, cols=1)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            cell = table.cell(0, 0)
            shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="EFF6FF"/>')
            cell._tc.get_or_add_tcPr().append(shading_elm)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            quote_text = stripped[2:].replace("*", "").replace('"', "")
            run = p.add_run(f'"{quote_text}"')
            run.italic = True
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
        elif stripped.startswith("- ") or stripped.startswith("* "):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            # Add runs supporting bold
            parts = re.split(r"(\*\*.*?\*\*)", stripped[2:])
            for part in parts:
                if part.startswith("**") and part.endswith("**"):
                    run = p.add_run(part[2:-2])
                    run.bold = True
                else:
                    p.add_run(part)
        elif re.match(r"^\d+\.\s", stripped):
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            m = re.match(r"^\d+\.\s(.*)", stripped)
            text_val = m.group(1) if m else stripped
            parts = re.split(r"(\*\*.*?\*\*)", text_val)
            for part in parts:
                if part.startswith("**") and part.endswith("**"):
                    run = p.add_run(part[2:-2])
                    run.bold = True
                else:
                    p.add_run(part)
        elif stripped == "---":
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(
                "__________________________________________________________________________"
            )
            run.font.color.rgb = RGBColor(0xCC, 0xD4, 0xDD)
        elif stripped:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            parts = re.split(r"(\*\*.*?\*\*)", stripped)
            for part in parts:
                if part.startswith("**") and part.endswith("**"):
                    run = p.add_run(part[2:-2])
                    run.bold = True
                else:
                    p.add_run(part)

        i += 1

    doc.save(str(output_path))
    print(f"[OK] Created Word document: {output_path}")


def create_html(md_path: Path, output_path: Path):
    content = md_path.read_text(encoding="utf-8")

    # Basic markdown to HTML conversion
    html_lines = []
    in_code = False
    in_list = False

    for line in content.splitlines():
        if line.startswith("```"):
            if in_code:
                html_lines.append("</code></pre>")
                in_code = False
            else:
                lang = line[3:].strip()
                html_lines.append(f'<pre><code class="language-{lang}">')
                in_code = True
            continue

        if in_code:
            escaped = line.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            html_lines.append(escaped)
            continue

        stripped = line.strip()

        if stripped.startswith("# "):
            html_lines.append(f"<h1>{stripped[2:]}</h1>")
        elif stripped.startswith("## "):
            html_lines.append(f"<h2>{stripped[3:]}</h2>")
        elif stripped.startswith("### "):
            html_lines.append(f"<h3>{stripped[4:]}</h3>")
        elif stripped.startswith("#### "):
            html_lines.append(f"<h4>{stripped[5:]}</h4>")
        elif stripped.startswith("> "):
            html_lines.append(f"<blockquote>{stripped[2:]}</blockquote>")
        elif stripped.startswith("- ") or stripped.startswith("* "):
            if not in_list:
                html_lines.append("<ul>")
                in_list = True
            item_text = re.sub(r"\*\*(.*?)\*\*", r"<strong>\1</strong>", stripped[2:])
            html_lines.append(f"<li>{item_text}</li>")
        elif re.match(r"^\d+\.\s", stripped):
            if not in_list:
                html_lines.append("<ol>")
                in_list = True
            item_text = re.sub(r"^\d+\.\s", "", stripped)
            item_text = re.sub(r"\*\*(.*?)\*\*", r"<strong>\1</strong>", item_text)
            html_lines.append(f"<li>{item_text}</li>")
        elif stripped == "---":
            if in_list:
                html_lines.append("</ul>")
                in_list = False
            html_lines.append("<hr/>")
        elif stripped:
            if in_list:
                html_lines.append("</ul>")
                in_list = False
            p_text = re.sub(r"\*\*(.*?)\*\*", r"<strong>\1</strong>", stripped)
            p_text = re.sub(r"`(.*?)`", r"<code>\1</code>", p_text)
            html_lines.append(f"<p>{p_text}</p>")
        else:
            if in_list:
                html_lines.append("</ul>")
                in_list = False

    if in_list:
        html_lines.append("</ul>")

    body_html = "\n".join(html_lines)

    html_template = f"""<!DOCTYPE html>
<html lang="en">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>NexoraERP & EKA — Interview Demonstration Playbook</title>
    <style>
        :root {{
            --primary: #2563eb;
            --primary-dark: #1e40af;
            --text-main: #1e293b;
            --text-muted: #64748b;
            --bg-main: #ffffff;
            --bg-alt: #f8fafc;
            --border: #e2e8f0;
            --code-bg: #0f172a;
            --code-text: #f8fafc;
        }}
        body {{
            font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, "Helvetica Neue", Arial, sans-serif;
            color: var(--text-main);
            background-color: var(--bg-alt);
            line-height: 1.65;
            margin: 0;
            padding: 40px 20px;
        }}
        .container {{
            max-width: 960px;
            margin: 0 auto;
            background: var(--bg-main);
            padding: 48px 64px;
            border-radius: 12px;
            box-shadow: 0 4px 20px -2px rgba(0, 0, 0, 0.08);
            border: 1px solid var(--border);
        }}
        h1 {{
            color: var(--primary-dark);
            font-size: 28px;
            border-bottom: 2px solid var(--border);
            padding-bottom: 12px;
            margin-top: 0;
        }}
        h2 {{
            color: var(--primary);
            font-size: 22px;
            margin-top: 32px;
            border-bottom: 1px solid var(--border);
            padding-bottom: 8px;
        }}
        h3 {{
            color: var(--text-main);
            font-size: 18px;
            margin-top: 24px;
        }}
        h4 {{
            color: var(--text-muted);
            font-size: 15px;
            margin-top: 18px;
            text-transform: uppercase;
            letter-spacing: 0.05em;
        }}
        blockquote {{
            background-color: #eff6ff;
            border-left: 4px solid var(--primary);
            margin: 16px 0;
            padding: 16px 20px;
            border-radius: 0 8px 8px 0;
            color: #1e3a8a;
            font-style: italic;
        }}
        pre {{
            background: var(--code-bg);
            color: var(--code-text);
            padding: 16px 20px;
            border-radius: 8px;
            overflow-x: auto;
            font-size: 13.5px;
            line-height: 1.5;
            margin: 16px 0;
        }}
        code {{
            font-family: Consolas, Monaco, "Courier New", monospace;
            background: #f1f5f9;
            color: #0f172a;
            padding: 2px 6px;
            border-radius: 4px;
            font-size: 13.5px;
        }}
        pre code {{
            background: transparent;
            color: inherit;
            padding: 0;
        }}
        ul, ol {{
            padding-left: 24px;
        }}
        li {{
            margin-bottom: 6px;
        }}
        hr {{
            border: none;
            border-top: 1px solid var(--border);
            margin: 32px 0;
        }}
        .print-btn {{
            float: right;
            background: var(--primary);
            color: white;
            border: none;
            padding: 8px 16px;
            border-radius: 6px;
            font-size: 14px;
            font-weight: 600;
            cursor: pointer;
            box-shadow: 0 2px 4px rgba(37, 99, 235, 0.2);
        }}
        .print-btn:hover {{
            background: var(--primary-dark);
        }}
        @media print {{
            body {{
                background: white;
                padding: 0;
            }}
            .container {{
                box-shadow: none;
                border: none;
                padding: 0;
                max-width: 100%;
            }}
            .print-btn {{
                display: none;
            }}
        }}
    </style>
</head>
<body>
    <div class="container">
        <button class="print-btn" onclick="window.print()">🖨️ Print / Save to PDF</button>
        {body_html}
    </div>
</body>
</html>"""

    output_path.write_text(html_template, encoding="utf-8")
    print(f"[OK] Created HTML document: {output_path}")


if __name__ == "__main__":
    root = Path(__file__).resolve().parent.parent
    md_file = root / "docs" / "INTERVIEW_DEMO_PLAYBOOK.md"
    docx_file = root / "docs" / "INTERVIEW_DEMO_PLAYBOOK.docx"
    html_file = root / "docs" / "INTERVIEW_DEMO_PLAYBOOK.html"

    create_docx(md_file, docx_file)
    create_html(md_file, html_file)
